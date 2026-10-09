"""File validation and text extraction for PDF, TXT and DOCX."""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from src.utils import (
    DocumentLoadError,
    EmptyDocumentError,
    RAGError,
    file_sha256,
    get_logger,
)

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def validate_file(path: Path, max_mb: int) -> None:
    """Check existence, extension, emptiness and size before parsing."""
    if not path.exists() or not path.is_file():
        raise DocumentLoadError(f"File not found: '{path.name}'.")
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentLoadError(
            f"'{path.name}' has unsupported type '{extension}'. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )
    size = path.stat().st_size
    if size == 0:
        raise EmptyDocumentError(f"'{path.name}' is empty (0 bytes).")
    if size > max_mb * 1024 * 1024:
        raise DocumentLoadError(f"'{path.name}' is larger than the {max_mb} MB limit.")


def _load_pdf(path: Path, file_hash: str) -> list[Document]:
    """One Document per page so page numbers survive into citations."""
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentLoadError(f"'{path.name}' is password-protected.")
        total = len(reader.pages)
        documents: list[Document] = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": path.name,
                            "file_type": "pdf",
                            "file_hash": file_hash,
                            "page": number,
                            "total_pages": total,
                        },
                    )
                )
    except RAGError:
        raise
    except Exception as exc:  # pypdf raises many different error types
        raise DocumentLoadError(f"Could not read PDF '{path.name}': {exc}") from exc
    if not documents:
        raise EmptyDocumentError(
            f"'{path.name}' contains no extractable text. It may be a scanned "
            "image PDF; this project does not include OCR."
        )
    return documents


def _load_txt(path: Path, file_hash: str) -> list[Document]:
    data = path.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    if not text.strip():
        raise EmptyDocumentError(f"'{path.name}' contains no text.")
    return [
        Document(
            page_content=text,
            metadata={"source": path.name, "file_type": "txt", "file_hash": file_hash},
        )
    ]


def _load_docx(path: Path, file_hash: str) -> list[Document]:
    try:
        document = DocxDocument(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
    except Exception as exc:
        raise DocumentLoadError(f"Could not read DOCX '{path.name}': {exc}") from exc
    text = "\n\n".join(parts)
    if not text.strip():
        raise EmptyDocumentError(f"'{path.name}' contains no text.")
    return [
        Document(
            page_content=text,
            metadata={"source": path.name, "file_type": "docx", "file_hash": file_hash},
        )
    ]


_LOADERS = {".pdf": _load_pdf, ".txt": _load_txt, ".docx": _load_docx}


def load_file(path: Path, max_mb: int = 50) -> list[Document]:
    """Validate and load one file into raw (uncleaned) Documents."""
    validate_file(path, max_mb)
    loader = _LOADERS[path.suffix.lower()]
    return loader(path, file_sha256(path))


def load_documents(paths: list[Path], max_mb: int = 50) -> tuple[list[Document], list[str]]:
    """Load many files. A bad file becomes a warning instead of aborting the batch."""
    documents: list[Document] = []
    warnings: list[str] = []
    for path in paths:
        try:
            loaded = load_file(path, max_mb)
            documents.extend(loaded)
            logger.info("Loaded %s (%d section(s))", path.name, len(loaded))
        except RAGError as exc:
            warnings.append(str(exc))
    return documents, warnings
